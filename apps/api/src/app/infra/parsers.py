import html as html_lib
import io
import re
from dataclasses import dataclass
from pathlib import PurePath

import trafilatura
from docx import Document as DocxDocument
from pypdf import PdfReader

_TITLE = re.compile(r"<title[^>]*>(.*?)</title>", re.IGNORECASE | re.DOTALL)


class UnsupportedFileError(ValueError):
    pass


@dataclass(frozen=True, slots=True)
class RawDocument:
    title: str
    text: str
    url: str | None = None


SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md", ".html", ".htm"}


def parse_html(html: str, url: str | None = None) -> RawDocument:
    text = trafilatura.extract(
        html,
        url=url,
        output_format="markdown",
        include_tables=True,
        include_links=False,
        include_comments=False,
        favor_recall=True,
    )
    match = _TITLE.search(html)
    title = (
        html_lib.unescape(re.sub(r"\s+", " ", match.group(1)).strip()) if match else (url or "Page")
    )
    return RawDocument(title=title, text=text or "", url=url)


def parse_pdf(data: bytes) -> str:
    reader = PdfReader(io.BytesIO(data))
    return "\n\n".join((page.extract_text() or "").strip() for page in reader.pages)


def parse_docx(data: bytes) -> str:
    doc = DocxDocument(io.BytesIO(data))
    parts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n\n".join(parts)


def decode_text(data: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1251"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def parse_file(filename: str, data: bytes) -> RawDocument:
    path = PurePath(filename)
    suffix = path.suffix.lower()
    title = path.stem.replace("_", " ").replace("-", " ").strip() or filename
    if suffix == ".pdf":
        return RawDocument(title=title, text=parse_pdf(data))
    if suffix == ".docx":
        return RawDocument(title=title, text=parse_docx(data))
    if suffix in {".txt", ".md"}:
        return RawDocument(title=title, text=decode_text(data))
    if suffix in {".html", ".htm"}:
        return parse_html(decode_text(data))
    raise UnsupportedFileError(f"Unsupported file type: {suffix or filename}")
