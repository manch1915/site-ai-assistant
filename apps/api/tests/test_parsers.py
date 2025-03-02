import io

import pytest
from docx import Document as DocxDocument

from app.infra.crawler import is_public_address
from app.infra.parsers import UnsupportedFileError, parse_file, parse_html
from tests.factories import minimal_pdf

HTML = """
<html><head><title>Prices | Ohm &amp; Solder</title></head>
<body>
  <nav>Home Prices Contact</nav>
  <main>
    <h1>Repair prices</h1>
    <p>iPhone 15 screen replacement costs 189 EUR and usually takes about an hour.</p>
    <p>Diagnostics are free if you go ahead with the repair. Otherwise it is 15 EUR.</p>
  </main>
  <footer>Copyright</footer>
</body></html>
"""


def test_parse_html_keeps_main_text_and_title() -> None:
    doc = parse_html(HTML, url="https://ohmsolder.test/prices")
    assert doc.title == "Prices | Ohm & Solder"
    assert "189 EUR" in doc.text
    assert doc.url == "https://ohmsolder.test/prices"


def test_parse_docx_reads_paragraphs_and_tables() -> None:
    document = DocxDocument()
    document.add_paragraph("Warranty terms")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Battery"
    table.rows[0].cells[1].text = "6 months"
    buffer = io.BytesIO()
    document.save(buffer)

    parsed = parse_file("warranty-terms.docx", buffer.getvalue())
    assert parsed.title == "warranty terms"
    assert "Warranty terms" in parsed.text
    assert "Battery | 6 months" in parsed.text


def test_parse_pdf_extracts_text() -> None:
    parsed = parse_file("price-list.pdf", minimal_pdf("Laptop keyboard replacement 79 EUR"))
    assert "Laptop keyboard replacement 79 EUR" in parsed.text


def test_parse_txt_falls_back_to_cp1251() -> None:
    parsed = parse_file("notes.txt", "Гарантия на ремонт".encode("cp1251"))
    assert parsed.text == "Гарантия на ремонт"


def test_unsupported_extension_is_rejected() -> None:
    with pytest.raises(UnsupportedFileError):
        parse_file("archive.zip", b"PK")


@pytest.mark.parametrize(
    ("address", "public"),
    [
        ("8.8.8.8", True),
        ("127.0.0.1", False),
        ("10.1.2.3", False),
        ("169.254.169.254", False),
        ("100.64.0.1", False),
        ("::1", False),
        ("::ffff:127.0.0.1", False),
        ("0.0.0.0", False),
        ("224.0.0.1", False),
    ],
)
def test_only_global_addresses_are_crawlable(address: str, public: bool) -> None:
    assert is_public_address(address) is public
